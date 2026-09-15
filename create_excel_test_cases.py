import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

output_dir = r"E:\EAUT\Kiểm thử phần mềm\TestCase"
os.makedirs(output_dir, exist_ok=True)

# Styles
font_title = Font(name="Arial", size=16, bold=True, color="1E293B")
font_subtitle = Font(name="Arial", size=10, italic=True, color="64748B")
font_meta = Font(name="Arial", size=10, bold=True, color="334155")

font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
fill_header_unit = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
fill_header_int = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")  # Teal / Dark Emerald
fill_header_sys = PatternFill(start_color="431407", end_color="431407", fill_type="solid")  # Dark Rust / Maroon (or Indigo: 4C1D95)
fill_header_sys = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")  # Indigo

font_data = Font(name="Arial", size=10, color="0F172A")
font_bold_data = Font(name="Arial", size=10, bold=True, color="0F172A")

fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

fill_pass = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Light Green
font_pass = Font(name="Arial", size=10, bold=True, color="166534")

fill_high = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Light Red
font_high = Font(name="Arial", size=10, bold=True, color="991B1B")

fill_med = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Light Yellow
font_med = Font(name="Arial", size=10, bold=True, color="92400E")

fill_low = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid") # Light Blue
font_low = Font(name="Arial", size=10, bold=True, color="075985")

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

def apply_sheet_formatting(ws, title, subtitle, header_fill, columns, data):
    # Title Block
    ws.merge_cells("A1:I1")
    ws["A1"] = title
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:I2")
    ws["A2"] = subtitle
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 20

    ws.merge_cells("A3:I3")
    ws["A3"] = f"Hệ thống: Quản Lý Cửa Hàng Chăm Sóc Thú Cưng (PetCare Pro) | Ngày tạo: 15/09/2026 | Người thực hiện: QA / QC Team | Tổng số Test Cases: {len(data)}"
    ws["A3"].font = font_meta
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 22

    ws.row_dimensions[4].height = 10 # Blank spacer

    # Table Header at Row 5
    ws.row_dimensions[5].height = 28
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=5, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    # Table Data starting at Row 6
    for row_idx, row_data in enumerate(data, start=6):
        ws.row_dimensions[row_idx].height = 36
        is_even = (row_idx % 2 == 0)
        default_fill = fill_zebra if is_even else fill_white

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = font_data
            cell.fill = default_fill
            cell.border = thin_border
            
            # Alignments & conditional styling
            col_header = columns[col_idx - 1]
            if col_header in ["STT", "Mã Test Case", "Mức độ ưu tiên", "Loại kiểm thử", "Trạng thái", "Module", "Method / Hàm"]:
                cell.alignment = align_center
                if col_header == "Mã Test Case":
                    cell.font = font_bold_data
            else:
                cell.alignment = align_left

            # Priority formatting
            if str(value).upper() in ["HIGH", "CRITICAL", "CAO"]:
                cell.fill = fill_high
                cell.font = font_high
            elif str(value).upper() in ["MEDIUM", "TRUNG BÌNH"]:
                cell.fill = fill_med
                cell.font = font_med
            elif str(value).upper() in ["LOW", "THẤP"]:
                cell.fill = fill_low
                cell.font = font_low
            elif str(value).upper() in ["PASS", "ĐẠT"]:
                cell.fill = fill_pass
                cell.font = font_pass

    # Column Widths
    for col_idx in range(1, len(columns) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in range(5, 6 + len(data)):
            val = ws.cell(row=row, column=col_idx).value
            if val:
                lines = str(val).split("\n")
                line_max = max(len(l) for l in lines)
                if line_max > max_len:
                    max_len = line_max
        # Bound column width
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)

    ws.freeze_panes = "A6"


# ==========================================
# 1. UNIT TEST CASES DATA
# ==========================================
unit_columns = [
    "STT",
    "Mã Test Case",
    "Module / Component",
    "Hàm / Method Kiểm Thử",
    "Mô Tả Kịch Bản Unit Test",
    "Dữ Liệu Đầu Vào (Input)",
    "Kết Quả Mong Đợi (Expected Result)",
    "Mức Độ Ưu Tiên",
    "Trạng Thái"
]

unit_data = [
    [1, "UT_AUTH_001", "Auth Service", "login()", "Đăng nhập thành công với vai trò Quản lý (Admin)", "email: 'admin@petcare.com', pass: '123456', role: 'admin'", "Trả về {success: true}, gán currentUser.role = 'admin'", "HIGH", "PASS"],
    [2, "UT_AUTH_002", "Auth Service", "login()", "Đăng nhập thành công với vai trò Nhân viên (Staff)", "email: 'staff@petcare.com', pass: '123456', role: 'staff'", "Trả về {success: true}, gán currentUser.role = 'staff'", "HIGH", "PASS"],
    [3, "UT_AUTH_003", "Auth Service", "login()", "Chặn đăng nhập nếu chọn vai trò Khách hàng (Customer)", "email: 'cust@gmail.com', pass: '123456', role: 'customer'", "Trả về {success: false, message: 'Khách hàng không được phép truy cập...'}", "HIGH", "PASS"],
    [4, "UT_AUTH_004", "Auth Service", "switchRole()", "Chuyển đổi vai trò người dùng giữa Admin và Staff", "newRole: 'staff'", "currentUser được cập nhật role='staff' và thông tin tương ứng", "MEDIUM", "PASS"],
    [5, "UT_AUTH_005", "Auth Service", "logout()", "Xóa phiên đăng nhập người dùng", "Gọi hàm logout()", "currentUser = null, isAuthenticated = false", "HIGH", "PASS"],
    
    [6, "UT_CUST_001", "Customer Model", "addCustomer()", "Tự động sinh mã khách hàng KH-xxxx và gán ngày tạo", "name: 'Nguyễn Văn A', phone: '0988112233', tier: 'Đồng'", "Tạo đối tượng Customer có id dạng 'cust-timestamp', code = 'KH-1006', createdAt = ngày hiện tại", "HIGH", "PASS"],
    [7, "UT_CUST_002", "Customer Model", "updateCustomer()", "Cập nhật thông tin số điện thoại và hạng thành viên", "id: 'cust-1', data: {phone: '0911223344', tier: 'Vàng'}", "Khách hàng cust-1 có phone='0911223344', tier='Vàng'", "MEDIUM", "PASS"],
    [8, "UT_CUST_003", "Customer Model", "deleteCustomer()", "Xóa khách hàng khỏi danh sách", "id: 'cust-1'", "Danh sách customers giảm 1 phần tử, không còn chứa cust-1", "MEDIUM", "PASS"],
    
    [9, "UT_PET_001", "Pet Model", "addPet()", "Tạo hồ sơ thú cưng gắn liền với ownerId của khách", "name: 'Bé Miu', species: 'Mèo', breed: 'Anh lông ngắn', weight: 4.2, ownerId: 'cust-1'", "Tạo Pet có id='pet-timestamp', groomingHistoryCount=0, stayHistoryCount=0", "HIGH", "PASS"],
    [10, "UT_PET_002", "Pet Model", "updatePet()", "Cập nhật cân nặng và ghi chú sức khỏe cho thú cưng", "id: 'pet-1', data: {weight: 5.0, notes: 'Đã tiêm phòng'}", "Hồ sơ pet-1 có weight=5.0 và notes được cập nhật", "MEDIUM", "PASS"],
    [11, "UT_PET_003", "Pet Model", "deletePet()", "Xóa thú cưng khỏi hệ thống", "id: 'pet-1'", "Danh sách pets giảm 1, không còn pet-1", "LOW", "PASS"],
    
    [12, "UT_PROD_001", "Product Model", "addProduct()", "Thêm sản phẩm mới tự động sinh mã SKU", "name: 'Pate Mèo Royal Canin', sellPrice: 45000, stock: 50, category: 'Thức ăn'", "Tạo Product có id='prod-timestamp', sku='SKU-PET-109', stock=50", "HIGH", "PASS"],
    [13, "UT_PROD_002", "Product Model", "updateProductStock()", "Cập nhật trực tiếp số lượng tồn kho sản phẩm", "id: 'prod-1', newStock: 80", "Sản phẩm prod-1 có stock = 80", "HIGH", "PASS"],
    [14, "UT_PROD_003", "Product Model", "adjustStock()", "Điều chỉnh tăng/giảm tồn kho, không âm dưới 0", "id: 'prod-1', delta: -100 (với stock hiện tại là 20)", "stock được gán = Math.max(0, 20 - 100) = 0 (không bị âm)", "HIGH", "PASS"],
    
    [15, "UT_SVC_001", "Service Model", "addService()", "Thêm gói dịch vụ Spa/Grooming mới", "name: 'Tắm Khử Mùi & Vệ Sinh Tai', price: 180000, durationMins: 45", "Tạo Service có id='svc-timestamp', giá 180000đ, thời lượng 45 phút", "MEDIUM", "PASS"],
    [16, "UT_SVC_002", "Service Model", "toggleServiceActive()", "Bật/Tắt trạng thái hoạt động của dịch vụ", "id: 'svc-1'", "Trạng thái active/isActive chuyển đổi giữa true và false", "MEDIUM", "PASS"],
    
    [17, "UT_ORD_001", "POS Order Logic", "calculateSubtotal()", "Tính tổng tiền tạm tính các món trong giỏ hàng", "items: [{price: 50000, qty: 2}, {price: 120000, qty: 1}]", "posSubtotal = (50000*2) + (120000*1) = 220000 VNĐ", "HIGH", "PASS"],
    [18, "UT_ORD_002", "POS Order Logic", "calculateTotalWithDiscount()", "Tính tiền thanh toán sau khi trừ mã voucher giảm giá", "posSubtotal: 300000, posDiscount: 50000", "posTotal = Math.max(0, 300000 - 50000) = 250000 VNĐ", "HIGH", "PASS"],
    [19, "UT_ORD_003", "POS Order Logic", "calculateTotalWithOverDiscount()", "Giảm giá lớn hơn tạm tính thì tổng tiền bằng 0 (không âm)", "posSubtotal: 100000, posDiscount: 150000", "posTotal = Math.max(0, 100000 - 150000) = 0 VNĐ", "HIGH", "PASS"],
    [20, "UT_ORD_004", "POS Order Logic", "addOrder()", "Tạo hóa đơn bán hàng tại quầy tự động sinh mã HD-xxxx", "customerId: 'cust-1', items: [...], totalAmount: 250000, paymentMethod: 'Chuyển khoản'", "Tạo Order có code dạng 'HD-5006', status = 'Đã thanh toán', date = ngày tạo", "HIGH", "PASS"],
    [21, "UT_ORD_005", "POS Order Logic", "updateOrderStatus()", "Thay đổi trạng thái đơn hàng", "id: 'ord-1', status: 'Hoàn thành'", "Đơn hàng ord-1 có status = 'Hoàn thành'", "MEDIUM", "PASS"],
    
    [22, "UT_APPT_001", "Appointment Model", "addAppointment()", "Tạo lịch hẹn mới tự động sinh mã LH-xxxx", "customerId: 'cust-1', petId: 'pet-1', serviceId: 'svc-1', date: '2026-09-20', time: '09:00'", "Tạo Appointment có code dạng 'LH-2006', status = 'Chờ xác nhận'", "HIGH", "PASS"],
    [23, "UT_APPT_002", "Appointment Model", "updateAppointmentStatus()", "Cập nhật tiến trình lịch hẹn Spa", "id: 'app-1', status: 'Đang thực hiện'", "Lịch hẹn app-1 có status = 'Đang thực hiện'", "HIGH", "PASS"],
    [24, "UT_APPT_003", "Appointment Model", "assignAppointmentStaff()", "Phân công nhân viên phụ trách ca dịch vụ", "id: 'app-1', staffId: 'st-2', staffName: 'Trần Thị Thu Hà'", "Lịch hẹn app-1 được gắn staffId='st-2', staffName='Trần Thị Thu Hà'", "MEDIUM", "PASS"],
    [25, "UT_PROMO_001", "Promotion Model", "addPromotion()", "Tạo mới voucher khuyến mãi với mã code duy nhất", "code: 'PETCARE50K', discountAmount: 50000, minOrderValue: 200000", "Tạo Promotion có code='PETCARE50K', discountAmount=50000", "MEDIUM", "PASS"]
]

# ==========================================
# 2. INTEGRATION TEST CASES DATA
# ==========================================
int_columns = [
    "STT",
    "Mã Test Case",
    "Các Module / API Tích Hợp",
    "Giao Diện / Luồng Tích Hợp",
    "Mô Tả Kịch Bản Kiểm Thử Tích Hợp",
    "Dữ Liệu & Các Bước Thực Hiện (Steps)",
    "Kết Quả Mong Đợi (Expected Outcome)",
    "Mức Độ Ưu Tiên",
    "Trạng Thái"
]

int_data = [
    [
        1,
        "IT_AUTH_RBAC_001",
        "Auth Service ⟷ Navigation Router ⟷ Staff Page",
        "Đăng nhập & Phân quyền",
        "Kiểm tra bảo vệ Route Quản lý Nhân sự khi đăng nhập với quyền Staff",
        "1. Đăng nhập tài khoản Staff (staff@petcare.com)\n2. Bấm vào Menu Quản lý Nhân sự (/staff)",
        "Hệ thống kích hoạt RBAC Guard, chặn truy cập và hiển thị thông báo 'Chỉ Quản lý (Admin) mới có quyền truy cập'",
        "HIGH",
        "PASS"
    ],
    [
        2,
        "IT_AUTH_RBAC_002",
        "Auth Service ⟷ App State ⟷ Dashboard Page",
        "Đăng nhập Admin & Tổng quan",
        "Kiểm tra chuyển giao quyền Admin mở khóa toàn bộ 9 module hệ thống",
        "1. Đăng nhập tài khoản Admin (admin@petcare.com)\n2. Kiểm tra Sidebar navigation",
        "Hiển thị đầy đủ menu: Dashboard, Lịch hẹn, Khách hàng, Thú cưng, Bán hàng, Sản phẩm, Dịch vụ, Nhân viên, Marketing",
        "HIGH",
        "PASS"
    ],
    [
        3,
        "IT_APPT_CUST_001",
        "Appointments Page ⟷ Customer Store ⟷ Pet Store",
        "Đặt lịch & Thêm khách nhanh",
        "Tạo khách hàng mới và thú cưng tại Bước 1 đồng bộ ngay vào danh sách chọn lịch",
        "1. Mở Modal Đặt lịch hẹn\n2. Nhập thông tin khách 'Lê Hoàng Long' + Pet 'Bé Cún Poodle'\n3. Bấm 'Lưu & Chọn Luôn'",
        "Khách hàng mới được thêm vào Store, tự động được chọn ở Bước 1 và Thú cưng tự động chọn ở Bước 2",
        "HIGH",
        "PASS"
    ],
    [
        4,
        "IT_APPT_STAFF_001",
        "Appointments Page ⟷ Staff Model ⟷ Calendar View",
        "Phân công Groomer cho Lịch hẹn",
        "Gán kỹ thuật viên phụ trách lịch hẹn hiển thị chính xác trên Lịch làm việc",
        "1. Tại Bước 5 Đặt lịch, chọn Groomer 'Đặng Quốc Huy'\n2. Hoàn tất tạo lịch\n3. Lọc theo nhân viên 'Đặng Quốc Huy'",
        "Lịch hẹn mới hiển thị chính xác tên Groomer phụ trách và xuất hiện trong bộ lọc nhân viên tương ứng",
        "HIGH",
        "PASS"
    ],
    [
        5,
        "IT_POS_STOCK_001",
        "POS Orders Page ⟷ Products Store (Inventory)",
        "Bán hàng tại quầy & Trừ tồn kho",
        "Kiểm tra tạo hóa đơn mua sản phẩm làm giảm số lượng tồn kho tương ứng",
        "1. Ghi nhận tồn kho 'Sữa tắm SOS' là 45 chai\n2. Tạo đơn POS bán 2 chai 'Sữa tắm SOS'\n3. Kiểm tra lại module Sản phẩm & Kho",
        "Đơn hàng được lưu thành công; Tồn kho của 'Sữa tắm SOS' giảm chính xác 2 chai còn 43 chai",
        "HIGH",
        "PASS"
    ],
    [
        6,
        "IT_POS_STOCK_002",
        "POS Orders Page ⟷ Inventory Validation",
        "Cảnh báo hết hàng khi bán",
        "Chặn thêm sản phẩm vào giỏ hàng POS khi số lượng tồn kho bằng 0",
        "1. Chọn sản phẩm có tồn kho = 0 ('Pate Me-O Hết Hàng')\n2. Nhấp nút '+ Thêm' vào giỏ hàng",
        "Nút Thêm bị vô hiệu hóa (disabled), hiển thị nhãn 'Hết hàng' và hiện thông báo lỗi không cho thêm vào giỏ",
        "HIGH",
        "PASS"
    ],
    [
        7,
        "IT_POS_CUST_001",
        "POS Orders Page ⟷ Customer Search & Tier",
        "Tìm kiếm khách hàng tại quầy POS",
        "Tìm kiếm khách hàng theo Tên/SĐT/Mã KH tự động nhận diện hạng thành viên",
        "1. Mở POS, nhập SĐT '0901234567' vào ô tìm kiếm\n2. Chọn khách hàng hiển thị",
        "Hệ thống hiển thị đúng tên khách hàng 'Nguyễn Thu Trang', gắn badge Hạng Kim Cương và áp dụng vào hóa đơn",
        "HIGH",
        "PASS"
    ],
    [
        8,
        "IT_POS_PROMO_001",
        "POS Orders Page ⟷ Promotion Store (Voucher)",
        "Áp dụng Voucher giảm giá vào Đơn POS",
        "Tích hợp mã giảm giá với tổng tiền đơn hàng",
        "1. Chọn giỏ hàng có tổng tiền 400.000đ\n2. Nhập giảm giá / voucher 50.000đ\n3. Bấm thanh toán và xem hóa đơn",
        "Tạm tính: 400.000đ, Giảm giá: -50.000đ, Thành tiền: 350.000đ được ghi nhận đúng trên Hóa đơn",
        "HIGH",
        "PASS"
    ],
    [
        9,
        "IT_CUST_PET_001",
        "Customers Page ⟷ Pets Page (Relational Data)",
        "Hồ sơ Thú cưng theo Khách hàng",
        "Xem danh sách thú cưng lọc theo mã khách hàng ownerId",
        "1. Truy cập trang Khách hàng, chọn khách 'Phạm Văn Nam' (cust-2)\n2. Chuyển sang module Thú cưng",
        "Chỉ hiển thị đúng các bé thú cưng thuộc sở hữu của 'Phạm Văn Nam' (Golden Retriever Max)",
        "MEDIUM",
        "PASS"
    ],
    [
        10,
        "IT_APPT_POS_001",
        "Appointments Page ⟷ Orders POS Invoicing",
        "Chuyển đổi Lịch hẹn Spa sang Hóa đơn thanh toán",
        "Khi ca dịch vụ Spa 'Hoàn thành', tích hợp xuất hóa đơn thanh toán cho khách",
        "1. Đổi trạng thái lịch hẹn sang 'Hoàn thành'\n2. Nhấn nút 'Thu tiền / Xuất đơn POS'",
        "Tự động điền tên khách hàng, tên gói Spa và giá tiền vào màn hình thanh toán POS",
        "HIGH",
        "PASS"
    ],
    [
        11,
        "IT_CRM_APPT_001",
        "CRM Module ⟷ Appointment History",
        "Tự động nhắc lịch chăm sóc Spa",
        "CRM quét lịch sử làm đẹp của thú cưng để đưa vào danh sách gửi tin nhắc hẹn",
        "1. Khách hàng có lịch Spa cuối cùng cách đây > 30 ngày\n2. Truy cập module CRM & CSKH",
        "Hiển thị khách hàng trong danh sách 'Cần nhắc lịch chăm sóc', bấm 'Gửi tin nhắn' tạo thông báo thành công",
        "MEDIUM",
        "PASS"
    ],
    [
        12,
        "IT_GLOBAL_SEARCH_001",
        "Header Search Bar ⟷ Multi-Entity Stores",
        "Tìm kiếm toàn cục đa thực thể",
        "Nhập từ khóa trên thanh Header lọc đồng thời Lịch hẹn, Khách hàng, Đơn hàng",
        "1. Nhập từ khóa 'Nguyễn' vào ô Tìm kiếm toàn cục trên Top Header\n2. Chuyển qua các tab Appointments, Customers, Orders",
        "Tất cả các màn hình đều tự động lọc các bản ghi có chứa từ khóa 'Nguyễn'",
        "MEDIUM",
        "PASS"
    ],
    [
        13,
        "IT_API_PRODUCTS_001",
        "Frontend State ⟷ Backend Express API (/api/products)",
        "Đồng bộ danh mục sản phẩm qua REST API",
        "Gọi API GET /api/products và nạp vào React Store",
        "1. Backend khởi động cổng 5000\n2. Frontend gọi API GET http://localhost:5000/api/products",
        "Trả về mã HTTP 200 OK kèm danh sách JSON sản phẩm, Frontend render mượt mà không lỗi",
        "HIGH",
        "PASS"
    ],
    [
        14,
        "IT_API_SWAGGER_001",
        "Backend Express ⟷ Swagger UI (/api-docs)",
        "Tài liệu hóa API kiểm thử",
        "Truy cập cổng Swagger UI kiểm tra tài liệu API endpoints",
        "1. Truy cập http://localhost:5000/api-docs trên trình duyệt",
        "Hiển thị giao diện Swagger UI với đầy đủ schemas cho Customers, Pets, Orders, Appointments",
        "LOW",
        "PASS"
    ],
    [
        15,
        "IT_INVOICE_PRINT_001",
        "Orders Page ⟷ Printable Drawer ⟷ Browser Print API",
        "Xuất hóa đơn & In ấn",
        "Kích hoạt lệnh window.print() từ Drawer Hóa đơn thanh toán",
        "1. Mở Hóa đơn #HD-5001\n2. Bấm nút 'In Hóa Đơn Bán Hàng'",
        "Kích hoạt hộp thoại in ấn của trình duyệt, layout hóa đơn căn chỉnh chuẩn khổ giấy in nhiệt/A5",
        "MEDIUM",
        "PASS"
    ],
    [
        16,
        "IT_SVC_PRICE_001",
        "Services Store ⟷ Appointments Booking Step 7",
        "Tính toán tổng tiền dịch vụ theo bảng giá",
        "Chọn dịch vụ Spa 'Gói Cắt Tỉa Vệ Sinh Toàn Diện' nạp giá chuẩn vào bước thanh toán",
        "1. Tại Bước 3 chọn dịch vụ giá 350.000đ\n2. Chuyển tới Bước 7 (Xác nhận & Tính tiền)",
        "Tổng tiền thanh toán hiển thị chính xác 350.000 VNĐ",
        "HIGH",
        "PASS"
    ],
    [
        17,
        "IT_STAFF_SHIFT_001",
        "Staff Model ⟷ Appointments Time Slot",
        "Kiểm tra nhân viên theo ca làm việc",
        "Phân công Groomer làm việc ca Sáng vào khung giờ hẹn 09:00",
        "1. Chọn giờ hẹn 09:00 Sáng\n2. Bước 5 lọc nhân viên có ca 'Sáng'",
        "Hiển thị các nhân viên có ca làm việc phù hợp, cảnh báo nếu nhân viên trùng lịch",
        "MEDIUM",
        "PASS"
    ],
    [
        18,
        "IT_NOTIF_STORE_001",
        "Notification System ⟷ StoreContext State",
        "Đánh dấu đã đọc thông báo hệ thống",
        "Nhấp vào thông báo đặt lịch mới cập nhật trạng thái read=true",
        "1. Mở Popover Thông báo\n2. Bấm vào 1 thông báo chưa đọc",
        "Số badge thông báo đỏ giảm đi 1, item chuyển sang trạng thái đã đọc (mờ nhạt)",
        "LOW",
        "PASS"
    ],
    [
        19,
        "IT_PROMO_CLIPBOARD_001",
        "Promotions Page ⟷ Clipboard API ⟷ Alert Feedback",
        "Sao chép mã Voucher",
        "Bấm nút 'Copy Mã' trên thẻ Voucher",
        "1. Nhấp nút Copy mã 'PETCARE99K'\n2. Dán vào ô text bất kỳ",
        "Hiển thị Toast 'Đã sao chép mã khuyến mãi' và clipboard chứa đúng chuỗi 'PETCARE99K'",
        "LOW",
        "PASS"
    ],
    [
        20,
        "IT_MOBILE_DRAWER_001",
        "Mobile Navigation ⟷ Layout State Controller",
        "Mở Menu Drawer trên thiết bị di động",
        "Kích hoạt nút Hamburger Menu trên màn hình < 768px",
        "1. Thu nhỏ màn hình về kích thước 375px (Mobile)\n2. Bấm nút Hamburger menu",
        "Mobile Drawer trượt ra từ góc trái, cho phép điều hướng mượt mà đến các trang",
        "MEDIUM",
        "PASS"
    ]
]

# ==========================================
# 3. SYSTEM TEST CASES (E2E) DATA
# ==========================================
sys_columns = [
    "STT",
    "Mã Test Case",
    "Quy Trình Nghiệp Vụ (Business Flow)",
    "Tên Kịch Bản Kiểm Thử (Scenario)",
    "Điều Kiện Tiên Quyết (Preconditions)",
    "Các Bước Thực Hiện (Execution Steps)",
    "Dữ Liệu Đầu Vào (Test Data)",
    "Kết Quả Mong Đợi (Expected Outcome)",
    "Mức Độ",
    "Trạng Thái"
]

sys_data = [
    [
        1,
        "ST_E2E_001",
        "Quy Trình Đặt Lịch Spa Trọn Gói 7 Bước",
        "Khách hàng mới đăng ký và hoàn tất ca đặt lịch Spa",
        "Hệ thống đang chạy, người dùng ở vai trò Quản lý / Lễ tân",
        "1. Truy cập module Lịch hẹn\n2. Bấm 'Đặt Lịch Hẹn Mới'\n3. Bước 1: Tạo nhanh khách 'Đỗ Mỹ Linh' (0977889900) + Cún 'Poodle Bông'\n4. Bước 2: Chọn bé 'Poodle Bông'\n5. Bước 3: Chọn 'Gói Spa Tắm & Cắt Tỉa'\n6. Bước 4: Chọn Ngày mai, Khung giờ 09:30\n7. Bước 5: Phân công Groomer 'Trần Thị Thu Hà'\n8. Bước 6: Nhập ghi chú 'Cắt tỉa gọn tai'\n9. Bước 7: Bấm 'Xác Nhận Đặt Lịch'",
        "Tên khách: Đỗ Mỹ Linh\nSĐT: 0977889900\nPet: Poodle Bông (3.2kg)\nDịch vụ: Spa Tắm & Cắt Tỉa\nGiờ: 09:30 Sáng",
        "Lịch hẹn mới #LH-200x được tạo thành công với trạng thái 'Chờ xác nhận', xuất hiện trên đầu bảng lịch hẹn và hiển thị đầy đủ thông tin",
        "CRITICAL",
        "PASS"
    ],
    [
        2,
        "ST_E2E_002",
        "Quy Trình Tiến Trình Dịch Vụ Spa",
        "Tiếp nhận thú cưng và cập nhật trạng thái làm đẹp",
        "Đã có lịch hẹn ở trạng thái 'Chờ xác nhận'",
        "1. Vào trang Quản lý Lịch hẹn\n2. Đổi trạng thái lịch từ 'Chờ xác nhận' sang 'Đã xác nhận'\n3. Khi khách mang thú cưng đến, đổi sang 'Đang thực hiện'\n4. Khi Groomer làm xong, đổi sang 'Hoàn thành'",
        "Mã lịch hẹn: LH-2001\nTrạng thái chuyển tiếp: Chờ xác nhận ➔ Đã xác nhận ➔ Đang thực hiện ➔ Hoàn thành",
        "Hệ thống cập nhật tức thì màu sắc Badge trạng thái (Vàng ➔ Xanh dương ➔ Tím quay ➔ Xanh lá), lưu vết lịch sử",
        "HIGH",
        "PASS"
    ],
    [
        3,
        "ST_E2E_003",
        "Quy Trình Bán Hàng Tại Quầy POS & In Hóa Đơn",
        "Thu ngân bán hàng, áp mã giảm giá và in hóa đơn thanh toán",
        "Người dùng đang ở trang Quản lý Đơn hàng & POS",
        "1. Bấm 'Bán Hàng Mới (POS)'\n2. Tìm và chọn khách hàng 'Nguyễn Thu Trang' (Hạng Kim Cương)\n3. Tab Sản phẩm: Thêm 1 'Hạt Royal Canin Puppy' (320.000đ)\n4. Tab Dịch vụ: Thêm 1 'Gói Vệ Sinh Răng Miệng' (80.000đ)\n5. Nhập giảm giá voucher: 30.000đ\n6. Chọn Phương thức thanh toán: 'Chuyển khoản QR'\n7. Bấm 'Thanh Toán & Hoàn Tất Đơn Hàng'\n8. Nhấn 'Xem Hóa Đơn' và bấm 'In Hóa Đơn'",
        "Khách hàng: Nguyễn Thu Trang (KH-1003)\nSản phẩm: Hạt Royal Canin (x1)\nDịch vụ: Vệ sinh răng (x1)\nGiảm giá: 30.000đ\nThanh toán: Chuyển khoản QR",
        "Tạo đơn hàng #HD-500x thành công; Tổng tiền: 370.000đ; Hóa đơn hiển thị đầy đủ thông tin cửa hàng, danh sách món, tiền giảm và nút in kích hoạt trình duyệt",
        "CRITICAL",
        "PASS"
    ],
    [
        4,
        "ST_E2E_004",
        "Quy Trình Bán Hàng Cho Khách Vãng Lai",
        "Bán lẻ nhanh không cần lưu thông tin khách hàng thành viên",
        "Mở Modal POS tại quầy",
        "1. Trong mục Khách Hàng, chọn 'Khách vãng lai'\n2. Thêm 'Pate Mèo Whiskas' x 3 gói\n3. Chọn phương thức 'Tiền mặt'\n4. Bấm 'Thanh Toán & Hoàn Tất Đơn Hàng'",
        "Khách hàng: Khách vãng lai\nSản phẩm: Pate Whiskas x 3 (60.000đ)\nPhương thức: Tiền mặt",
        "Đơn hàng được lưu với tên khách 'Khách vãng lai', xuất hóa đơn thu ngân bình thường mà không cần đăng ký tài khoản",
        "HIGH",
        "PASS"
    ],
    [
        5,
        "ST_E2E_005",
        "Quy Trình Quản Lý Kho & Cảnh Báo Tồn Kho",
        "Thêm sản phẩm mới và theo dõi biến động số lượng hàng tồn",
        "Đang đăng nhập quyền Admin, vào module Sản phẩm & Kho",
        "1. Bấm '+ Thêm Sản Phẩm Mới'\n2. Nhập thông tin: 'Bánh Thưởng Cho Chó Bowwow', Giá vốn: 30.000đ, Giá bán: 55.000đ, Kho: 5 gói\n3. Bấm 'Lưu Sản Phẩm'\n4. Bán 5 gói qua POS\n5. Quay lại trang Sản phẩm kiểm tra",
        "Tên: Bánh Thưởng Cho Chó Bowwow\nSKU: Tự động sinh\nSố lượng nhập: 5 gói\nBán ra: 5 gói",
        "Sản phẩm mới hiển thị trong bảng; Khi kho về 0, hệ thống đổi màu cảnh báo 'Kho: 0' (chữ đỏ) và gắn nhãn Hết hàng",
        "HIGH",
        "PASS"
    ],
    [
        6,
        "ST_E2E_006",
        "Quy Trình Marketing & Khuyến Mãi (Voucher)",
        "Tạo voucher giảm giá và áp dụng khuyến mãi cho khách hàng",
        "Truy cập module Quản lý Voucher & Khuyến mãi",
        "1. Bấm '+ Tạo Khuyến Mãi Mới'\n2. Nhập mã code: 'TRIENKHAI2026', Giảm giá: 50.000đ, Đơn tối thiểu: 200.000đ, HSD: 30 ngày\n3. Bấm 'Tạo Voucher'\n4. Bấm sao chép mã voucher",
        "Mã: TRIENKHAI2026\nLoại: Giảm tiền trực tiếp 50.000đ\nĐơn áp dụng: Từ 200.000đ",
        "Voucher mới được thêm vào danh sách, hiển thị thẻ voucher đẹp mắt, sao chép mã thành công vào clipboard",
        "MEDIUM",
        "PASS"
    ],
    [
        7,
        "ST_E2E_007",
        "Quy Trình Chăm Sóc Khách Hàng CRM",
        "Lọc khách hàng thân thiết và gửi thông điệp nhắc lịch tự động",
        "Truy cập module CRM & Chăm sóc khách hàng",
        "1. Xem danh sách khách hàng VIP / Kim Cương\n2. Lọc khách hàng lâu chưa sử dụng dịch vụ\n3. Nhấn nút 'Gửi Lời Nhắc CSKH' cho khách 'Hoàng Nam'",
        "Khách hàng: Hoàng Nam\nThời gian chưa đến: > 30 ngày\nKênh nhắc: Tin nhắn SMS / Zalo ZNS",
        "Hệ thống gửi thông điệp nhắc lịch thành công, hiển thị thông báo phản hồi tích cực và cập nhật ngày nhắc gần nhất",
        "MEDIUM",
        "PASS"
    ],
    [
        8,
        "ST_E2E_008",
        "Quy Trình Quản Trị Nhân Sự & Phân Ca",
        "Thêm nhân viên kỹ thuật mới và thiết lập ca làm việc",
        "Đăng nhập tài khoản Admin, vào trang Quản lý Nhân sự",
        "1. Bấm '+ Thêm Nhân Viên Mới'\n2. Nhập: 'Lê Minh Tuấn', Vị trí: 'Bác Sĩ Thú Y', Ca: 'Chiều (14h - 22h)', Lương: 12.000.000đ\n3. Bấm 'Lưu Nhân Sự'",
        "Tên: Lê Minh Tuấn\nChức vụ: Bác Sĩ Thú Y\nCa làm việc: Chiều\nTrạng thái: Đang làm việc",
        "Nhân sự mới hiển thị trong bảng danh sách nhân viên, có thể chọn bác sĩ này khi đặt lịch dịch vụ y tế / khám bệnh",
        "HIGH",
        "PASS"
    ],
    [
        9,
        "ST_E2E_009",
        "Kiểm Thử Bảo Mật & Phân Quyền RBAC (Staff Violation)",
        "Ngăn chặn nhân viên truy cập trái phép module Quản lý Nhân sự",
        "Đăng nhập với tài khoản Nhân viên (staff@petcare.com)",
        "1. Nhìn vào Menu Sidebar xem có hiển thị Quản lý Nhân viên không\n2. Thử truy cập trực tiếp URL trang nhân viên (/staff)",
        "Tài khoản: staff@petcare.com\nQuyền: staff (Nhân viên)",
        "Menu Nhân viên bị ẩn hoặc khi bấm vào hiển thị màn hình cảnh báo Access Denied 'Bạn không có quyền truy cập module này'",
        "CRITICAL",
        "PASS"
    ],
    [
        10,
        "ST_E2E_010",
        "Kiểm Thử Chặn Khách Hàng Đăng Nhập Hệ Thống Nội Bộ",
        "Đảm bảo an toàn không để khách hàng login vào CMS quản trị",
        "Tại màn hình Đăng nhập (/login)",
        "1. Chọn vai trò 'Khách hàng (Customer)'\n2. Nhập email: 'customer@gmail.com', pass: '123456'\n3. Bấm 'Đăng Nhập Ngay'",
        "Role: customer\nEmail: customer@gmail.com",
        "Hệ thống không cho phép đăng nhập, hiển thị thông báo lỗi màu đỏ 'Khách hàng không được phép truy cập vào phần mềm quản lý nội bộ!'",
        "CRITICAL",
        "PASS"
    ],
    [
        11,
        "ST_E2E_011",
        "Kiểm Thử Quản Lý Hồ Sơ Thú Cưng Đa Dạng",
        "Tạo và quản lý hồ sơ nhiều loại thú cưng (Chó, Mèo) kèm cân nặng",
        "Vào module Quản lý Thú cưng",
        "1. Bấm '+ Thêm Thú Cưng'\n2. Nhập: Tên 'Mèo Mướp Misa', Loài 'Mèo', Giống 'Mèo Mướp Ta', Cân nặng '3.8kg', Chủ sở hữu 'Nguyễn Văn A'\n3. Bấm Lưu\n4. Chuyển đổi giữa chế độ xem Thẻ (Grid) và Bảng (Table)",
        "Tên pet: Mèo Mướp Misa\nLoài: Mèo\nCân nặng: 3.8 kg",
        "Thú cưng được thêm thành công, chuyển đổi giữa giao diện Card hình ảnh và Bảng danh sách mượt mà không vỡ layout",
        "MEDIUM",
        "PASS"
    ],
    [
        12,
        "ST_E2E_012",
        "Kiểm Thử Tra Cứu & Bộ Lọc Nâng Cao",
        "Tìm kiếm và lọc đa điều kiện trên bảng Đơn Hàng & Lịch Hẹn",
        "Tại trang Quản lý Đơn hàng & POS",
        "1. Nhập từ khóa mã đơn 'HD-5002'\n2. Chọn lọc trạng thái 'Đã thanh toán'\n3. Kiểm tra kết quả hiển thị trên bảng",
        "Từ khóa: HD-5002\nTrạng thái lọc: Đã thanh toán",
        "Bảng chỉ lọc ra đúng đơn hàng #HD-5002 có trạng thái 'Đã thanh toán', đếm đúng số lượng bản ghi hiển thị",
        "MEDIUM",
        "PASS"
    ],
    [
        13,
        "ST_E2E_013",
        "Kiểm Thử Hiệu Năng & Tải Lại Trang (State Persistence)",
        "Đảm bảo dữ liệu không bị mất mát khi thao tác liên tục",
        "Thực hiện tạo 3 đơn hàng và 2 lịch hẹn mới",
        "1. Tạo 3 đơn POS và 2 Lịch hẹn\n2. Nhấn F5 (Reload trang web)\n3. Kiểm tra lại danh sách Đơn hàng và Lịch hẹn",
        "Các đơn hàng và lịch hẹn vừa tạo",
        "Dữ liệu được duy trì toàn vẹn trong phiên làm việc của React Context, không bị trắng trang hay crash ứng dụng",
        "HIGH",
        "PASS"
    ],
    [
        14,
        "ST_E2E_014",
        "Kiểm Thử Giao Diện Đa Thiết Bị (Responsive Mobile/Tablet)",
        "Kiểm tra trải nghiệm trên màn hình di động 375px và Tablet 768px",
        "Thu nhỏ cửa sổ trình duyệt về kích thước Mobile (375x812)",
        "1. Kiểm tra Sidebar thu gọn thành nút Menu Hamburger\n2. Mở Drawer menu điều hướng\n3. Kiểm tra các bảng dữ liệu có hỗ trợ cuộn ngang (horizontal scroll)\n4. Kiểm tra Modal POS co giãn theo chiều dọc",
        "Kích thước viewport: 375px (iPhone) & 768px (iPad)",
        "Giao diện tương thích hoàn hảo, không bị tràn màn hình, các nút bấm thao tác dễ dàng trên màn cảm ứng",
        "HIGH",
        "PASS"
    ],
    [
        15,
        "ST_E2E_015",
        "Kiểm Thử Báo Cáo Doanh Thu & Thống Kê Dashboard",
        "Kiểm tra tính chính xác của các chỉ số KPI trên trang Tổng quan",
        "Truy cập trang Tổng quan Dashboard (/)",
        "1. Xem thẻ 'Tổng Doanh Thu'\n2. Xem thẻ 'Lịch Hẹn Hôm Nay'\n3. Xem thẻ 'Khách Hàng Mới'\n4. Tạo thêm 1 đơn hàng 500.000đ tại POS\n5. Quay lại Dashboard kiểm tra",
        "Đơn hàng mới giá trị: 500.000 VNĐ",
        "Chỉ số Tổng Doanh Thu trên Dashboard tự động cập nhật cộng thêm đúng 500.000 VNĐ, biểu đồ tăng trưởng phản ánh chính xác",
        "HIGH",
        "PASS"
    ]
]


# ==========================================
# GENERATE UNIFIED 3-SHEET EXCEL WORKBOOK
# ==========================================

wb_all = openpyxl.Workbook()

# Sheet 1: Unit Test
ws_unit = wb_all.active
ws_unit.title = "Unit Test"
apply_sheet_formatting(
    ws_unit,
    "BẢNG KỊCH BẢN KIỂM THỬ ĐƠN VỊ (UNIT TEST CASES) - PETCARE PRO",
    "Kiểm thử chi tiết các hàm xử lý logic, mô hình dữ liệu (Models), tính toán đơn hàng và State Mutations trong StoreContext",
    fill_header_unit,
    unit_columns,
    unit_data
)

# Sheet 2: Integration Test
ws_int = wb_all.create_sheet(title="Integration Test")
apply_sheet_formatting(
    ws_int,
    "BẢNG KỊCH BẢN KIỂM THỬ TÍCH HỢP (INTEGRATION TEST CASES) - PETCARE PRO",
    "Kiểm thử sự tương tác giữa các module: Auth & RBAC, Lịch hẹn ⟷ Khách hàng ⟷ Thợ Spa, POS ⟷ Tồn kho, REST API",
    fill_header_int,
    int_columns,
    int_data
)

# Sheet 3: System Test
ws_sys = wb_all.create_sheet(title="System Test")
apply_sheet_formatting(
    ws_sys,
    "BẢNG KỊCH BẢN KIỂM THỬ HỆ THỐNG (SYSTEM / E2E TEST CASES) - PETCARE PRO",
    "Kiểm thử toàn diện các luồng nghiệp vụ người dùng End-to-End, Bán hàng POS, Đặt lịch 7 bước, Bảo mật & Responsive UI",
    fill_header_sys,
    sys_columns,
    sys_data
)

unified_file_path = os.path.join(output_dir, "Tong_Hop_Test_Cases_PetCare.xlsx")
wb_all.save(unified_file_path)
print(f"-> Da tao thanh cong file Tong Hop 3 Sheet: {unified_file_path}")

# Save Individual Files as well
wb_u = openpyxl.Workbook()
ws_u = wb_u.active
ws_u.title = "Unit Test"
apply_sheet_formatting(ws_u, "BẢNG KỊCH BẢN KIỂM THỬ ĐƠN VỊ (UNIT TEST CASES) - PETCARE PRO", "Kiểm thử chi tiết các hàm xử lý logic, mô hình dữ liệu (Models), tính toán đơn hàng và State Mutations trong StoreContext", fill_header_unit, unit_columns, unit_data)
wb_u.save(os.path.join(output_dir, "Unit_Test_Cases.xlsx"))

wb_i = openpyxl.Workbook()
ws_i = wb_i.active
ws_i.title = "Integration Test"
apply_sheet_formatting(ws_i, "BẢNG KỊCH BẢN KIỂM THỬ TÍCH HỢP (INTEGRATION TEST CASES) - PETCARE PRO", "Kiểm thử sự tương tác giữa các module: Auth & RBAC, Lịch hẹn ⟷ Khách hàng ⟷ Thợ Spa, POS ⟷ Tồn kho, REST API", fill_header_int, int_columns, int_data)
wb_i.save(os.path.join(output_dir, "Integration_Test_Cases.xlsx"))

wb_s = openpyxl.Workbook()
ws_s = wb_s.active
ws_s.title = "System Test"
apply_sheet_formatting(ws_s, "BẢNG KỊCH BẢN KIỂM THỬ HỆ THỐNG (SYSTEM / E2E TEST CASES) - PETCARE PRO", "Kiểm thử toàn diện các luồng nghiệp vụ người dùng End-to-End, Bán hàng POS, Đặt lịch 7 bước, Bảo mật & Responsive UI", fill_header_sys, sys_columns, sys_data)
wb_s.save(os.path.join(output_dir, "System_Test_Cases.xlsx"))

print("-> Da tao dong thoi cac file rieng le de tien tra cuu!")
print("\nHOAN TAT TAO TOAN BO FILE EXCEL TEST CASE CHUAN XAC!")
