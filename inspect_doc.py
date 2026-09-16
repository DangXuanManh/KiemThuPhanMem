import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document(r'E:\EAUT\Kiểm thử phần mềm\BaoCao\DeTai2_Nhom6.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    print(f"P{i:03d} [{p.style.name}]: {p.text}")
