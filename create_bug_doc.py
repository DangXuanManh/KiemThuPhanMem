import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("QUY TRÌNH & SƠ ĐỒ QUẢN LÝ LỖI PHẦN MỀM\n(DEFECT MANAGEMENT & BUG LIFE CYCLE)")
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0xC2, 0x41, 0x0C) # Orange 700

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle_p.add_run("Dự án: PetCare Store Management System | Môn học: Kiểm thử phần mềm")
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_paragraph() # Spacing

    # Section 1: Introduction
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("1. TỔNG QUAN VỀ QUY TRÌNH QUẢN LÝ LỖI (DEFECT MANAGEMENT PROCESS)")
    h1_run.font.color.rgb = RGBColor(0xEA, 0x58, 0x0C)
    
    p = doc.add_paragraph(
        "Quy trình quản lý lỗi (Bug Management Process) là một tập hợp các bước chuẩn hóa nhằm phát hiện, ghi nhận, "
        "phân loại, theo dõi và sửa chữa các sai sót (defects/bugs) phát sinh trong suốt vòng đời phát triển phần mềm. "
        "Mục tiêu chính là đảm bảo mọi lỗi đều được xử lý triệt để trước khi sản phẩm được bàn giao tới khách hàng."
    )

    # Highlight box
    table_box = doc.add_table(rows=1, cols=1)
    table_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table_box.cell(0, 0)
    set_cell_background(cell, "FFF7ED") # Light Orange
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    box_p = cell.paragraphs[0]
    box_run = box_p.add_run(
        "Nguyên tắc cốt lõi: Tất cả lỗi tìm thấy trong quá trình Kiểm thử tự động (Automated Testing) hoặc Kiểm thử thủ công "
        "(Manual Testing) đều phải được tạo Bug Ticket hợp lệ, gắn nhãn Mức độ nghiêm trọng (Severity) và Mức độ ưu tiên (Priority) rõ ràng."
    )
    box_run.font.size = Pt(10.5)
    box_run.font.bold = True
    box_run.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)

    doc.add_paragraph()

    # Section 2: Bug Life Cycle Flowchart
    h2 = doc.add_heading(level=1)
    h2_run = h2.add_run("2. SƠ ĐỒ VÒNG ĐỜI LỖI (BUG LIFE CYCLE DIAGRAM)")
    h2_run.font.color.rgb = RGBColor(0xEA, 0x58, 0x0C)

    doc.add_paragraph("Dưới đây là sơ đồ chi tiết chuyển dịch trạng thái của Lỗi từ lúc phát hiện cho tới khi đóng Ticket:")

    # Text-based Flowchart Table
    flow_table = doc.add_table(rows=9, cols=3)
    flow_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["Trạng Thái (Status)", "Mô Tả Hành Động", "Người Thao Tác (Actor)"]
    hdr_row = flow_table.rows[0]
    for i, h_text in enumerate(headers):
        cell = hdr_row.cells[i]
        set_cell_background(cell, "EA580C")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    stages = [
        ("1. NEW (Mới phát hiện)", "Tester phát hiện hành vi bất thường của ứng dụng và lập Bug Report mới.", "Tester"),
        ("2. ASSIGNED (Đã phân công)", "Test Lead / PM duyệt lỗi và phân công cho Developer phụ trách module.", "Test Lead / PM"),
        ("3. OPEN / FIXING (Đang sửa)", "Developer xác nhận lỗi hợp lệ và tiến hành phân tích, sửa code.", "Developer"),
        ("4. FIXED (Đã sửa xong)", "Developer đã sửa xong lỗi trên môi trường Test và yêu cầu Re-test.", "Developer"),
        ("5. PENDING RETEST (Chờ test lại)", "Lỗi đã được đẩy lên build mới, chờ Tester kiểm tra lại.", "QA / Tester"),
        ("6. VERIFIED / PASSED (Đã xác minh)", "Tester kiểm tra lại trên build mới và xác nhận lỗi không còn xuất hiện.", "Tester"),
        ("7. CLOSED (Đã đóng lỗi)", "Lỗi đã được giải quyết hoàn toàn và đóng Ticket thành công.", "Tester / Lead"),
        ("8. REOPENED (Tái mở lỗi)", "Tester kiểm tra lại nhưng lỗi vẫn còn xuất hiện ➔ Chuyển lại về Developer.", "Tester")
    ]

    for row_idx, data in enumerate(stages, start=1):
        row_cells = flow_table.rows[row_idx].cells
        bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row_cells[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    doc.add_paragraph()

    # Visual Flow Chart representation
    doc.add_heading(level=2, text="2.1 Sơ Đồ Khối Chuyển Trạng Thái (Visual State Flow)")
    
    diagram_table = doc.add_table(rows=1, cols=1)
    diagram_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    diag_cell = diagram_table.cell(0, 0)
    set_cell_background(diag_cell, "F1F5F9")
    set_cell_margins(diag_cell, top=150, bottom=150, left=200, right=200)
    
    diag_text = (
        "  [ NEW (Loi moi) ]\n"
        "           │\n"
        "           ▼\n"
        "  [ ASSIGNED (Phan cong Dev) ]\n"
        "           │\n"
        "           ├─────────────────────────┐\n"
        "           ▼                         ▼\n"
        "  [ OPEN / FIXING ]          [ REJECTED / DUPLICATE ]\n"
        "           │                         │\n"
        "           ▼                         ▼\n"
        "  [ FIXED (Da sua) ]          [ CLOSED (Tu choi/Dong) ]\n"
        "           │\n"
        "           ▼\n"
        "  [ PENDING RETEST ]\n"
        "           │\n"
        "      ┌────┴────────────────────────┐\n"
        "      ▼                             ▼\n"
        "[ PASS: Verified ]           [ FAIL: Reopened ]\n"
        "      │                             │\n"
        "      ▼                             └──────► (Quay lai OPEN/FIXING)\n"
        "[ CLOSED (Da dong) ]"
    )
    dp = diag_cell.paragraphs[0]
    dr = dp.add_run(diag_text)
    dr.font.name = 'Courier New'
    dr.font.size = Pt(9.5)
    dr.font.bold = True
    dr.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_paragraph()

    # Section 3: Severity & Priority
    h3 = doc.add_heading(level=1)
    h3_run = h3.add_run("3. PHÂN LOẠI MỨC ĐỘ NGHIÊM TRỌNG (SEVERITY) & ƯU TIÊN (PRIORITY)")
    h3_run.font.color.rgb = RGBColor(0xEA, 0x58, 0x0C)

    # Severity Table
    doc.add_heading(level=2, text="3.1 Phân Loại Mức Độ Nghiêm Trọng (Severity)")
    sev_table = doc.add_table(rows=5, cols=3)
    sev_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    s_headers = ["Mức Độ (Severity)", "Mô Tả Chi Tiết", "Ví Dụ Trong PetCare Store"]
    for i, h_text in enumerate(s_headers):
        cell = sev_table.rows[0].cells[i]
        set_cell_background(cell, "1E293B")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    sev_data = [
        ("S1 - Critical (Nghiêm trọng)", "Hệ thống treo, sập database, rò rỉ dữ liệu hoặc không thể đặt lịch.", "Server backend bị crash khi khách hàng ấn Đặt Lịch Hẹn."),
        ("S2 - High (Cao)", "Tính năng chính bị lỗi nhưng có phương án thay thế tạm thời.", "Khách hàng không thể thêm thú cưng mới từ màn hình cá nhân."),
        ("S3 - Medium (Vừa)", "Tính năng phụ bị sai logic hoặc giao diện bị lệch ảnh hưởng trải nghiệm.", "Giá tiền dịch vụ hiển thị sai định dạng phân cách hàng nghìn."),
        ("S4 - Low (Thấp)", "Lỗi chính tả, màu sắc nút chưa đúng thiết kế hoặc lỗi nhỏ không ảnh hưởng.", "Sai lỗi chính tả từ 'Grooming' thành 'Groming' trên trang chủ.")
    ]

    for row_idx, data in enumerate(sev_data, start=1):
        cells = sev_table.rows[row_idx].cells
        bg = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            if col_idx == 0:
                r.font.bold = True
                if "S1" in text: r.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
                elif "S2" in text: r.font.color.rgb = RGBColor(0xEA, 0x58, 0x0C)
                elif "S3" in text: r.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)

    doc.add_paragraph()

    # Section 4: Sample Bug Report Template
    h4 = doc.add_heading(level=1)
    h4_run = h4.add_run("4. MẪU PHIẾU BÁO CÁO LỖI THỰC TẾ (SAMPLE BUG REPORT)")
    h4_run.font.color.rgb = RGBColor(0xEA, 0x58, 0x0C)

    doc.add_paragraph("Dưới đây là một Bug Ticket hoàn chỉnh mẫu được ghi nhận trong dự án PetCare Store:")

    bug_table = doc.add_table(rows=8, cols=2)
    bug_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    b_data = [
        ("Bug ID / Title", "BUG-PET-042: Khách hàng không thể đặt lịch hẹn khi bỏ trống ghi chú"),
        ("Module / Feature", "Đặt Lịch Dịch Vụ (Booking Module)"),
        ("Detected By & Date", "Tester Nguyễn Văn A | 08/09/2026"),
        ("Severity & Priority", "Severity: S2 (High) | Priority: P1 (Immediate)"),
        ("Environment", "Windows 11, Chrome v128, Node.js v24.x, Server Localhost:5000"),
        ("Steps to Reproduce", "1. Đăng nhập tài khoản Customer\n2. Vào trang Đặt Lịch Hẹn -> Chọn dịch vụ Spa Poodle\n3. Chọn thú cưng & ngày giờ hẹn hợp lệ\n4. Để trống ô 'Ghi Chú Đặc Biệt' -> Nhấn 'Xác Nhận Đặt Lịch'"),
        ("Expected Result", "Hệ thống chấp nhận đặt lịch thành công với ghi chú trống."),
        ("Actual Result", "Hệ thống báo lỗi 'Internal Server Error 500' do Backend xử lý NULL ở trường notes.")
    ]

    for row_idx, (label, val) in enumerate(b_data):
        cells = bug_table.rows[row_idx].cells
        set_cell_background(cells[0], "FFF7ED")
        set_cell_background(cells[1], "FFFFFF")
        set_cell_margins(cells[0], top=100, bottom=100, left=120, right=120)
        set_cell_margins(cells[1], top=100, bottom=100, left=120, right=120)
        
        p0 = cells[0].paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)

        p1 = cells[1].paragraphs[0]
        r1 = p1.add_run(val)
        if "BUG-PET" in val:
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)

    doc.add_paragraph()

    # Section 5: Roles & Responsibilities
    h5 = doc.add_heading(level=1)
    h5_run = h5.add_run("5. TRÁCH NHIỆM & QUY TRÌNH PHỐI HỢP DỰ ÁN")
    h5_run.font.color.rgb = RGBColor(0xEA, 0x58, 0x0C)

    doc.add_paragraph(
        "- Tester / QA: Có trách nhiệm tái hiện lỗi (reproduce), cung cấp đầy đủ log và các bước thực hiện, kiểm tra lại (re-test) khi Developer bàn giao build mới.\n"
        "- Developer: Có trách nhiệm tiếp nhận ticket, phân tích nguyên nhân gốc rễ (Root Cause), sửa mã nguồn và viết Unit Test bổ sung để tránh lỗi lặp lại.\n"
        "- Test Lead / Project Manager: Theo dõi tiến độ giải quyết ticket, đánh giá các lỗi bị trì hoãn (Deferred) và phê duyệt đóng release."
    )

    # Save document
    output_path = "E:/Kiểm thử phần mềm/So_Do_Quan_Ly_Loi_PetCareStore.docx"
    doc.save(output_path)
    print("SUCCESS: Word document created successfully!")

if __name__ == "__main__":
    create_document()
